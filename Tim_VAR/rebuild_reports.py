from pathlib import Path
import json, calendar
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
PACKAGE=ROOT;R=PACKAGE/'results';F=PACKAGE/'figures';F.mkdir(exist_ok=True)
DOCS=PACKAGE/'reports';DOCS.mkdir(exist_ok=True)
meta=json.loads((R/'metadata.json').read_text())
acc=pd.read_csv(R/'accuracy.csv');diag=pd.read_csv(R/'diagnostics.csv')
g=pd.read_csv(R/'granger_tests.csv');st=pd.read_csv(R/'stationarity.csv')
ds=pd.read_csv(R/'descriptive_statistics_training.csv').set_index('variable')
sea=pd.read_csv(R/'seasonality_training.csv');ic=pd.read_csv(R/'lag_selection.csv')
dat=pd.read_csv(R/'analysis_data.csv',parse_dates=['date'])
main=acc[acc['case']=='BIC'].set_index('model');sens=acc[acc['case']=='AIC'].set_index('model')
plt.rcParams.update({'font.family':'DejaVu Serif','font.size':9,'axes.spines.top':False,
 'axes.spines.right':False,'axes.labelsize':9,'axes.titlesize':10,'legend.fontsize':8,
 'axes.edgecolor':'#666666','axes.labelcolor':'#222222','xtick.color':'#333333',
 'ytick.color':'#333333','savefig.facecolor':'white'})
colors=['#254D73','#2C7F7A','#AE7442','#835685']
def finish(fig,name):
    fig.savefig(F/(name+'.png'),dpi=220,bbox_inches='tight')
    fig.savefig(F/(name+'.svg'),bbox_inches='tight')
    plt.close(fig)

fig,axs=plt.subplots(1,2,figsize=(6.5,2.1),layout='constrained')
for ax,criterion,color in zip(axs,['BIC','AIC'],colors):
    vals=ic[criterion];p=int(ic.loc[vals.idxmin(),'p'])
    ax.plot(ic.p,vals,color=color,marker='o',markersize=2.5,lw=1.4)
    ax.scatter([p],[vals.min()],color='#B04637',s=27,zorder=5)
    ax.annotate(f'{p} lags',(p,vals.min()),xytext=(8,14),textcoords='offset points',fontsize=9,
                 arrowprops={'arrowstyle':'-','color':'#555555'})
    ax.set(title=criterion,xlabel='Lag order',ylabel='Criterion value',xticks=[1,4,9,15,20])
    ax.grid(axis='y',color='#DDDDDD',lw=.5)
finish(fig,'A1_lag_selection')

f=pd.read_csv(R/'forecasts_AIC.csv',parse_dates=['date'])
fig,ax=plt.subplots(figsize=(6.5,2.4),layout='constrained')
for label,color in zip(['Oil and gold','Gas and Treasury','All commodities','Full VAR'],colors):
    gain=np.cumsum((f.actual-f.AR)**2-(f.actual-f[label])**2)
    ax.plot(f.date,gain,label=label,color=color,lw=1.15)
ax.axhline(0,color='#777777',lw=.65)
ax.set(ylabel='Cumulative gain',xlabel='Forecast target date')
ax.xaxis.set_major_locator(mdates.MonthLocator(bymonth=[1,7]))
ax.xaxis.set_major_formatter(mdates.DateFormatter('%b %Y'))
ax.legend(ncol=4,loc='upper center',bbox_to_anchor=(.5,1.22),frameon=False)
ax.grid(axis='y',color='#DDDDDD',lw=.5)
finish(fig,'A2_cumulative_gain_AIC')

fig,ax=plt.subplots(figsize=(6.5,1.6),layout='constrained')
ax.plot(dat.date,dat.ES,color=colors[0],lw=.55)
ax.axvline(pd.Timestamp('2023-07-12'),color='#B04637',ls='--',lw=1)
ax.text(pd.Timestamp('2023-09-01'),-5.55,'Test period',fontsize=8,color='#87362A')
ax.set(xlabel='Observation date',ylabel='Equity log RV5')
ax.grid(axis='y',color='#DDDDDD',lw=.5)
finish(fig,'B1_equity_series')

fig,axs=plt.subplots(1,2,figsize=(6.5,2.05),layout='constrained')
for ax,s,title,color in zip(axs,['NG','C'],['Natural gas','Corn'],colors):
    a=sea[sea.variable==s]
    ax.plot(a.month,a.relative_to_overall,color=color,marker='o',markersize=3,lw=1.5)
    ax.axhline(1,color='#777777',ls='--',lw=.6)
    ax.set(title=title,xlabel='Month',ylabel='Mean RV5 / training mean',xticks=[1,3,5,7,9,12])
    ax.grid(axis='y',color='#DDDDDD',lw=.5)
finish(fig,'B2_seasonality')

def newdoc(title,appendix=False):
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11 if appendix else 12)
    normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=1.1
    normal.paragraph_format.space_after=Pt(6)
    for name in ['Title','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(14 if name=='Title' else 12);st.font.bold=True
        st.paragraph_format.space_before=Pt(8 if name!='Title' else 0)
        st.paragraph_format.space_after=Pt(6)
    d.core_properties.author='Tim Schmitt'
    d.core_properties.title=title
    d.add_paragraph(title,'Title')
    return d

def para(d,text,style=None):
    p=d.add_paragraph(text,style);p.paragraph_format.widow_control=True
    return p

def note(d,text):
    p=para(d,text);p.paragraph_format.space_after=Pt(6)
    for r in p.runs:r.font.size=Pt(9.5)
    return p

def caption(d,text):
    p=para(d,text);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5)
    p.paragraph_format.space_after=Pt(4)
    for r in p.runs:r.font.size=Pt(10);r.bold=True

def table(d,headers,rows,widths):
    t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for ir,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        if ir==0:
            rep=OxmlElement('w:tblHeader');trPr.append(rep)
        for j,(c,w) in enumerate(zip(row.cells,widths)):
            c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side,val in [('top','65'),('bottom','65'),('left','75'),('right','75')]:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),val);x.set(qn('w:type'),'dxa');mar.append(x)
            tcpr.append(mar)
            if ir==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF1');tcpr.append(shade)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=ir==0
    return t

def figure(d,name,width=6.45):
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(3)
    run=p.add_run();run.add_picture(str(F/(name+'.png')),width=Inches(width))
    # Accessible image description.
    run._r.xpath('.//wp:docPr')[0].set('descr',name.replace('_',' '))

def save(d,name):
    # Remove inherited template title borders and theme font overrides.
    for tree in [d.element,d.styles.element]:
        for border in list(tree.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    for style_name in ['Normal','Title','Heading 1','Heading 2']:
        style=d.styles[style_name]
        rpr=style.element.get_or_add_rPr()
        rf=rpr.find(qn('w:rFonts'))
        if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
        for key in list(rf.attrib):
            if 'theme' in key.lower():del rf.attrib[key]
        for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
    for p in d.paragraphs:
        if p.style.name in ['Title','Heading 1','Heading 2']:
            for r in p.runs:
                r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
    p=DOCS/(name+'.docx');d.save(p)
    # Plain text companion for integration into the group document.
    lines=[]
    for block in d.element.body:
        if block.tag==qn('w:p'):
            text=''.join(block.itertext())
            # python-docx's paragraph.text avoids XML duplicate text traversal.
            from docx.text.paragraph import Paragraph
            text=Paragraph(block,d).text
            if text:lines.append(text)
        elif block.tag==qn('w:tbl'):
            from docx.table import Table
            tt=Table(block,d)
            lines.append('\n'.join(' | '.join(c.text.replace('\n',' ') for c in row.cells) for row in tt.rows))
    (DOCS/(name+'.txt')).write_text('\n\n'.join(lines)+'\n')
    print('Saved',p)

doc=newdoc('Vector autoregression')
para(doc,'To identify the contribution of other markets, we compare VAR forecasts with an equity-only AR containing the same number of lags and calendar controls. This comparison separates cross-market information from the predictive value of equity volatility’s own history. A persistence forecast, equal to the previous retained observation, provides an additional benchmark.')
para(doc,'The dependent variable is log RV5 of E-mini S&P 500 futures. We estimate four VAR specifications: equity with oil and gold; equity with natural gas and the Treasury proxy; equity with all four commodities, including corn; and a full system containing all six series. Each equation is estimated by OLS with an intercept, 11 month indicators and four weekday indicators. The calendar controls accommodate recurring seasonal differences without pre-filtering the data. Treasury variation enters as the supplied squared yield-change proxy. Appendix B describes the variables and common sample.')
para(doc,'The initial training sample contains 3,127 observations through 12 July 2023. Lag selection uses only this sample: BIC selects four lags and AIC selects nine from a maximum of 20 in the full system. Four lags form the main specification; nine provide a sensitivity check. We hold the lag order fixed across models within each comparison and re-estimate on an expanding sample, producing 783 one-step forecasts from 13 July 2023 to 31 August 2026. Forecasts use earlier retained observations and assume that all inputs dated before the target are available. Appendix A reports the specification checks.')
caption(doc,'Table 1 Out-of-sample accuracy with four lags')
rows=[]
for nm in ['Persistence','AR','Oil and gold','Gas and Treasury','All commodities','Full VAR']:
    a=main.loc[nm];rows.append([nm if nm in ['Persistence','AR','Full VAR'] else 'VAR: '+nm.lower(),f'{a.MSE:.5f}',f'{a.MAE:.5f}',f'{-a.gain_vs_AR_pct:+.2f}' if nm!='AR' else '0.00'])
table(doc,['Model','MSE','MAE','MSE difference\nvs AR (%)'],rows,[2.65,1.05,1.05,1.75])
note(doc,'Notes: 783 common forecast targets; errors are measured in equity log RV5. All AR/VAR models include identical calendar controls. Source: own calculations using the group dataset.')
para(doc,'The equity AR reduces MSE by 18.14% relative to persistence. None of the four-lag VARs improves on the matched AR in either MSE or MAE. The full VAR’s MSE is 0.70% higher, while adding oil and gold raises it by 0.33%. The improvement over persistence therefore comes from modelling equity’s own dynamics and calendar effects; it cannot be attributed to the additional markets.')
para(doc,'The nine-lag oil-and-gold VAR achieves a small 0.14% MSE reduction relative to AR(9), but its MAE is higher. Broader systems again fail to improve MSE. First differencing and an additional Treasury delay also leave the AR–VAR comparison unchanged (Appendix A). Mixed stationarity evidence and remaining residual autocorrelation warrant caution with conventional Granger-test inference. Overall, additional markets provide little incremental forecast value for this equity benchmark over the selected evaluation period. This finding does not establish that spillovers are absent or that the same result holds for other equity markets or forecast horizons.').paragraph_format.keep_together=True
save(doc,'VAR_main_report')

doc=newdoc('Appendix A VAR specification and checks',True)
para(doc,'All systems include equity log RV5. Additional futures variables are their log RV5 series; Treasury variation is the supplied squared yield change in basis points squared. Every equation contains an intercept and 15 calendar indicators. The equity-only AR uses the same deterministic controls. Models are estimated by OLS, equation by equation.')
caption(doc,'Table A1 Initial training specifications with four lags')
dg=diag[diag['case']=='BIC'].set_index('model')
labels={'AR':'ES','Oil and gold':'ES, CL, GC','Gas and Treasury':'ES, NG, TY','All commodities':'ES, CL, GC, NG, C','Full VAR':'ES, CL, GC, NG, C, TY'}
rows=[]
for nm in labels:
    a=dg.loc[nm];rows.append([nm,labels[nm],int(a.parameters_per_equation),f'{a.max_root:.3f}','<0.001' if a.BG20_p<.001 else f'{a.BG20_p:.3f}'])
table(doc,['Model','Variables','Coefficients\nper equation','Largest\nroot','Serial-correlation\np-value'],rows,[1.32,1.73,1.1,.85,1.5])
note(doc,'ES: equity; CL: oil; GC: gold; NG: natural gas; C: corn; TY: Treasury proxy. Coefficient counts include the intercept and calendar indicators. Roots are companion-matrix eigenvalue moduli. Serial-correlation checks are Breusch–Godfrey LM tests with 20 residual lags, applied to the equity equation.')
para(doc,'All initial training estimates satisfy the stability condition, with roots inside the unit circle. However, the residual tests reject absence of serial correlation. Increasing to nine lags reduces but does not eliminate this problem: the full VAR’s residual-test p-value is 0.010. The models are therefore evaluated primarily on out-of-sample forecasting, with conventional in-sample test results interpreted cautiously.')
caption(doc,'Figure A1 Lag selection in the initial training sample')
figure(doc,'A1_lag_selection')
note(doc,'AIC and BIC are computed for the full six-variable system using the same 3,107 regression observations at every candidate order. BIC selects four lags and AIC nine. After selection, estimation uses the available sample for that order: 3,123 initial regression rows at four lags and 3,118 at nine. No test-period outcomes enter lag selection.')
para(doc,'Forecast sensitivity and in-sample predictability','Heading 1').paragraph_format.page_break_before=True
caption(doc,'Table A2 Forecast accuracy and joint Granger tests')
rows=[]
for nm in labels:
    if nm=='AR':p4=p9='—'
    else:
        p4=f'{g[(g["case"]=="BIC")&(g.model==nm)].p_value.iloc[0]:.3f}'
        p9=f'{g[(g["case"]=="AIC")&(g.model==nm)].p_value.iloc[0]:.3f}'
    rows.append([nm,f'{main.loc[nm,"MSE"]:.5f}',f'{sens.loc[nm,"MSE"]:.5f}',p4,p9])
table(doc,['Model','MSE\n4 lags','MSE\n9 lags','Granger p\n4 lags','Granger p\n9 lags'],rows,[2.1,1.1,1.1,1.15,1.05])
note(doc,'MSE uses the same 783 test targets. Granger tests use only the initial training sample and jointly set all additional-market lag coefficients in the equity equation to zero, conditional on equity lags and calendar controls. Reported p-values are conventional F-test values; residual serial correlation limits their reliability.')
para(doc,'At four lags, the oil-and-gold and all-commodity systems yield nominal p-values below 0.05. Neither finding persists at nine lags, and neither system improves four-lag forecast accuracy. These results illustrate why apparent in-sample predictability is insufficient to establish a forecasting gain. They do not identify structural causal effects.')
caption(doc,'Figure A2 Cumulative forecast gain over AR with nine lags')
figure(doc,'A2_cumulative_gain_AIC')
note(doc,'Cumulative gain equals the sum of AR squared errors minus VAR squared errors. Positive values favour the VAR. All models use nine lags and the same controls. The small final oil-and-gold advantage varies over the evaluation period.')
para(doc,'Two further checks retain all 783 targets. First differencing, with four lags selected by training-sample BIC, gives level-forecast MSEs of 0.37291 for the AR and 0.37312 for the full VAR. Adding one further retained-observation delay to the Treasury proxy gives 0.35669 for a sample-matched AR and 0.35915 for the full VAR. Both checks leave the main AR–VAR ranking unchanged.')
save(doc,'VAR_appendix')

doc=newdoc('Appendix B Data and preliminary checks',True)
para(doc,'The supplied group dataset contains 3,910 dated observations and 24 columns. Dates are unique and ordered, and all numeric entries are finite and non-missing. The VAR analysis retains every row and selects five existing futures log-RV5 columns plus the Treasury variation proxy. Stored log columns equal the natural logarithm of their corresponding RV5 columns; no additional log transformation is applied.')
caption(doc,'Table B1 Common observation sample')
table(doc,['Sample','Dates','Observations'],[
 ['Full panel','4 January 2011–31 August 2026','3,910'],
 ['Initial training','4 January 2011–12 July 2023','3,127'],
 ['Test targets','13 July 2023–31 August 2026','783']],[1.5,3.8,1.2])
note(doc,'The existing training cutoff is retained. Counts precede model-specific lag losses. The expanding procedure incorporates each observed test outcome only into subsequent forecasts.')
caption(doc,'Table B2 Variables and initial training descriptives')
marketlabels={'ES':'Equity','CL':'WTI oil','GC':'Gold','NG':'Natural gas','C':'Corn','TY':'Treasury proxy'}
rows=[]
for s in marketlabels:
    rows.append([marketlabels[s],'log_rv5_'+s if s!='TY' else 'rv_yield_10',f'{ds.loc[s,"mean"]:.3f}',f'{ds.loc[s,"std"]:.3f}'])
table(doc,['Series','CSV column','Mean','Std. deviation'],rows,[1.5,2.5,1.2,1.3])
note(doc,'Futures variables are natural log RV5. Treasury variation is in basis points squared, unlogged; 308 full-sample observations are zero. Units differ across these two representations.')
para(doc,'The Treasury proxy equals the squared basis-point change in the yield between successive retained dates, wherever the change can be verified within the file. It can therefore span a gap in the common calendar and is not intraday bond realised variance. The baseline uses its ordinary model lags and assumes previous-date information is available; an extra-delay check appears in Appendix A. No values are interpolated or seasonally pre-filtered.')
caption(doc,'Figure B1 Equity log realised variance and the evaluation period')
figure(doc,'B1_equity_series')
note(doc,'The dashed line marks the fixed training cutoff. The plot displays the full supplied panel for description; preliminary tests use only the initial training sample.')
para(doc,'Stationarity and seasonality','Heading 1').paragraph_format.page_break_before=True
caption(doc,'Table B3 Stationarity tests and persistence in the initial training sample')
rows=[]
for _,a in st.iterrows():rows.append([marketlabels[a.variable],f'{a.ADF:.3f}',f'{a.KPSS:.3f}',f'{ds.loc[a.variable,"ACF1"]:.3f}'])
table(doc,['Series','ADF statistic','KPSS statistic','Lag-one ACF'],rows,[2.05,1.5,1.5,1.45])
note(doc,'ADF: intercept, no trend, 20 lagged differences; 5% critical value approximately −2.86. KPSS: level-stationarity null and 28 Bartlett lags; 5% critical value 0.463. Tests use the series as entered, before calendar controls. ACF denotes sample autocorrelation.')
para(doc,'ADF rejects a unit root for every series, while KPSS rejects constant-mean level stationarity for every series. This mixed evidence prevents an unqualified claim that the data are stationary. Strong persistence, changing volatility regimes and deterministic patterns may contribute to the disagreement. Log levels are retained as the baseline forecasting specification, while the first-difference check in Appendix A assesses whether the AR–VAR comparison depends on this choice.')
caption(doc,'Figure B2 Monthly volatility patterns in natural gas and corn')
figure(doc,'B2_seasonality')
corn=sea[sea.variable=='C'].set_index('month')
ratio=float(corn.loc[7,'mean_RV5']/corn.loc[12,'mean_RV5'])
note(doc,'Each point is the arithmetic mean of unlogged RV5 for that month, divided by the series’ overall initial-training mean. These are descriptive monthly averages, not causal estimates or formal tests of seasonality.')
para(doc,f'Natural gas has its highest monthly mean in January and its lowest in August. Corn peaks in June; July’s mean RV5 is {ratio:.2f} times December’s. Rejecting a unit root does not remove these calendar patterns. The AR and VAR equations therefore include 11 month and four weekday indicators, estimated within each expanding training sample. The underlying observations remain unchanged, and seasonal coefficients are never estimated using future test outcomes.')
save(doc,'Data_appendix')
print('Main prose words:',sum(len(p.text.split()) for p in Document(DOCS/'VAR_main_report.docx').paragraphs))
